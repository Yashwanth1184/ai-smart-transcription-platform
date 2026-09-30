import json, os
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT
from icalendar import Calendar, Event
from datetime import datetime

def flatten(value, level=0):
    lines=[]
    if isinstance(value, dict):
        for k,v in value.items():
            lines.append((level, k.replace('_',' ').title()))
            lines += flatten(v, level+1)
    elif isinstance(value, list):
        for x in value: lines += flatten(x, level+1)
    else:
        lines.append((level, str(value)))
    return lines

def export_docx(note: dict, path: str):
    doc=Document(); doc.add_heading("AI Smart Notes",0)
    for level,text in flatten(note):
        if level==0: doc.add_heading(text,1)
        elif level==1: doc.add_paragraph(text, style="List Bullet" if isinstance(note.get(text), list) else None)
        else: doc.add_paragraph(text)
    doc.save(path)

def export_pdf(note: dict, path: str):
    styles=getSampleStyleSheet(); story=[Paragraph("AI Smart Notes", styles["Title"]), Spacer(1,12)]
    for level,text in flatten(note):
        style=styles["Heading2"] if level==0 else styles["BodyText"]
        story.append(Paragraph(text.replace('&','&amp;'), style)); story.append(Spacer(1,6))
    SimpleDocTemplate(path,pagesize=A4).build(story)

def export_txt(note: dict, path: str):
    with open(path,"w",encoding="utf-8") as f:
        for level,text in flatten(note): f.write("  "*level + text + "\n")

def export_ics(title: str, remind_at: str, path: str):
    cal=Calendar(); cal.add("prodid","-//AI Smart Notes//EN"); cal.add("version","2.0")
    event=Event(); event.add("summary",title); event.add("dtstart",datetime.fromisoformat(remind_at)); event.add("dtend",datetime.fromisoformat(remind_at)); cal.add_component(event)
    with open(path,"wb") as f:f.write(cal.to_ical())
